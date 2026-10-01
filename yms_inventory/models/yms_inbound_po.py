# -*- coding: utf-8 -*-

import base64
import csv
import io
import json

from odoo import _, api, fields, models
from odoo.exceptions import UserError


class YmsInboundPo(models.Model):
    """Inventory-side inbound PO reference. Not Purchase and not a stock ledger."""

    _name = "yms.inbound.po"
    _description = "YMS Inbound PO Reference"
    _order = "id desc"

    name = fields.Char(string="PO Number", required=True, index=True)
    supplier = fields.Char(string="Supplier")
    po_date = fields.Char(string="PO Date")
    expected_delivery = fields.Char(string="Expected Delivery")
    plant = fields.Char(string="Plant", default="Plant 1")
    status = fields.Char(string="Status", default="Open")
    remarks = fields.Char(string="Remarks", default="-")
    material = fields.Char(string="Material Summary")
    item_count = fields.Integer(string="Total Items", default=1)
    ordered_qty = fields.Float(string="Total Quantity")
    received_qty = fields.Float(string="Received So Far")
    pending_qty = fields.Float(string="Pending Quantity")
    arrival_datetime = fields.Char(string="Arrival Date/Time")
    truck_no = fields.Char(string="Truck No.")
    transporter = fields.Char(string="Transporter")
    reference_no = fields.Char(string="Reference No.")
    notes = fields.Char(string="Notes")
    invoice_no = fields.Char(string="Invoice / Challan")
    materials_json = fields.Text(string="Materials JSON")
    import_file = fields.Char(string="Last Import File")

    _sql_constraints = [("yms_inbound_po_name_uniq", "unique(name)", "PO Number must be unique.")]

    def _as_dict(self):
        self.ensure_one()
        try:
            materials = json.loads(self.materials_json or "[]")
        except json.JSONDecodeError:
            materials = []
        return {
            "id": self.id,
            "po": self.name,
            "supplier": self.supplier or "",
            "date": self.po_date or "",
            "eta": self.expected_delivery or "",
            "plant": self.plant or "Plant 1",
            "status": self.status or "Open",
            "remarks": self.remarks or "-",
            "material": self.material or "",
            "items": self.item_count,
            "ordered": self.ordered_qty,
            "received": self.received_qty,
            "pending": self.pending_qty,
            "arrival": {
                "datetime": self.arrival_datetime or "",
                "truck": self.truck_no or "",
                "transporter": self.transporter or "",
                "reference": self.reference_no or "",
                "notes": self.notes or "",
            },
            "invoice_no": self.invoice_no or "",
            "materials": materials,
            "import_file": self.import_file or "",
        }

    @api.model
    def yms_get_inbound_pos(self, query=None, status=None):
        domain = []
        if status and status != "all":
            domain.append(("status", "=", status))
        recs = self.search(domain, limit=200)
        rows = [rec._as_dict() for rec in recs]
        q = (query or "").strip().lower()
        if q:
            rows = [
                row
                for row in rows
                if q in f"{row['po']} {row['supplier']} {row['material']}".lower()
            ]
        return rows

    @api.model
    def yms_get_import_history(self):
        recs = self.search([("import_file", "!=", False)], limit=50)
        return [
            {
                "file": rec.import_file,
                "user": rec.write_uid.name if rec.write_uid else "",
                "time": fields.Datetime.to_string(rec.write_date) if rec.write_date else "",
                "rows": rec.item_count,
                "status": "Success",
            }
            for rec in recs
        ]

    @api.model
    def yms_import_inbound_csv(self, filename, content):
        return self.yms_import_inbound_file(filename, content)

    @api.model
    def yms_import_inbound_file(self, filename, content):
        name = str(filename or "").lower()
        payload = self._yms_decode_upload(content)
        if name.endswith(".xlsx"):
            po_rows, yard_rows = self._yms_xlsx_rows(payload)
        elif name.endswith(".csv"):
            po_rows = self._yms_csv_rows(payload)
            yard_rows = []
        else:
            raise UserError(_("Upload an Excel .xlsx or CSV file."))
        po_rows = [self._yms_alias_row(row) for row in po_rows]
        if po_rows and not po_rows[0].get("po") and po_rows[0].get("pipe_id"):
            raise UserError(_("This looks like the Yard Allocation file. Import YMS_Inbound_Input.xlsx on PO Import."))
        result = self._yms_import_po_rows(po_rows, filename)
        return result

    @api.model
    def yms_clear_inbound_pos(self):
        recs = self.search([])
        count = len(recs)
        recs.unlink()
        return {"cleared": count, "rows": [], "history": []}

    def _yms_decode_upload(self, content):
        raw = content or ""
        if isinstance(raw, bytes):
            return raw
        if str(raw).startswith("data:"):
            raw = str(raw).split(",", 1)[-1]
        try:
            return base64.b64decode(raw)
        except Exception as err:
            raise UserError(_("Could not read the file: %s") % err) from err

    def _yms_csv_rows(self, payload):
        try:
            text = payload.decode("utf-8-sig")
        except Exception as err:
            raise UserError(_("Could not read the CSV: %s") % err) from err
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise UserError(_("The file has no header row."))
        return [
            {k.strip().lower().replace(" ", "_"): (v or "").strip() for k, v in line.items() if k}
            for line in reader
        ]

    def _yms_xlsx_rows(self, payload):
        try:
            from openpyxl import load_workbook
        except ImportError as err:
            raise UserError(_("Excel import needs openpyxl on the Odoo server.")) from err
        try:
            workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
        except Exception as err:
            raise UserError(_("Could not read the Excel file: %s") % err) from err
        names = workbook.sheetnames
        skip = {"demo_guide", "how_to", "readme", "yard_plan", "yard_pipes"}
        if "Inbound_Receiving" in names:
            po_title = "Inbound_Receiving"
        elif "Inbound_PO" in names:
            po_title = "Inbound_PO"
        else:
            po_title = next((n for n in names if n.lower().replace(" ", "_") not in skip), names[0])
        po_rows = self._yms_sheet_dicts(workbook[po_title])
        yard_rows = []
        if "Yard_Plan" in names:
            yard_rows = self._yms_sheet_dicts(workbook["Yard_Plan"])
        elif "Yard_Pipes" in names:
            yard_rows = self._yms_sheet_dicts(workbook["Yard_Pipes"])
        return po_rows, yard_rows

    def _yms_yard_sheet_rows(self, payload):
        try:
            from openpyxl import load_workbook
        except ImportError as err:
            raise UserError(_("Excel import needs openpyxl on the Odoo server.")) from err
        workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
        names = workbook.sheetnames
        skip = {"demo_guide", "how_to", "readme", "inbound_receiving", "inbound_po"}
        if "Yard_Plan" in names:
            title = "Yard_Plan"
        elif "Yard_Pipes" in names:
            title = "Yard_Pipes"
        else:
            title = next((n for n in names if n.lower().replace(" ", "_") not in skip), names[0])
        return [self._yms_alias_row(row) for row in self._yms_sheet_dicts(workbook[title])]

    def _yms_yard_change_sheet_rows(self, payload):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return []
        workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
        names = workbook.sheetnames
        title = next((n for n in names if n.lower().replace(" ", "_") in ("location_change", "from_to", "yard_moves")), None)
        if not title:
            return []
        return [self._yms_alias_row(row) for row in self._yms_sheet_dicts(workbook[title])]

    def _yms_sheet_dicts(self, sheet):
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip().lower().replace(" ", "_") for h in rows[0]]
        out = []
        for values in rows[1:]:
            mapped = {}
            empty = True
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = values[idx] if idx < len(values) else None
                if value is None:
                    mapped[header] = ""
                else:
                    mapped[header] = str(value).strip()
                    if mapped[header]:
                        empty = False
            if not empty:
                out.append(self._yms_alias_row(mapped))
        return out

    def _yms_alias_row(self, mapped):
        aliases = {
            "po_no": "po",
            "po_number": "po",
            "purchase_order": "po",
            "grn_no": "grn_no",
            "grn_number": "grn_no",
            "material_description": "material_desc",
            "quantity": "ordered",
            "qty": "ordered",
            "invoice_challan": "invoice_no",
            "challan": "invoice_no",
            "challan_no": "invoice_no",
            "invoice": "invoice_no",
            "truck": "truck_no",
            "truck_number": "truck_no",
            "vehicle_no": "truck_no",
            "gate_ref": "reference_no",
            "reference": "reference_no",
            "arrival_date": "arrival_datetime",
            "arrival_time": "arrival_datetime",
            "inbound_status": "status",
            "expected_delivery_date": "expected_delivery",
            "specification": "spec",
            "heat": "batch",
            "heat_no": "batch",
            "batch_heat": "batch",
            "batch_no": "batch",
            "length_m": "length",
            "diameter_in": "diameter",
            "dia": "diameter",
            "to_yard": "yard",
            "to_bin": "bin",
            "to_plate": "plate",
            "from": "from_location",
            "from_loc": "from_location",
            "movement_type": "move_type",
            "inventory_change": "demo_note",
        }
        out = {}
        for key, value in (mapped or {}).items():
            norm = aliases.get(key, key)
            out[norm] = value
        if out.get("status") in ("Arrived", "Received"):
            out["status"] = "Open"
        if out.get("arrival_datetime"):
            out.setdefault("po_date", str(out["arrival_datetime"])[:10])
            out.setdefault("expected_delivery", str(out["arrival_datetime"])[:10])
        return out

    def _yms_import_po_rows(self, lines, filename):
        required = {"po", "supplier", "material_code", "ordered"}
        if not lines:
            raise UserError(_("No inbound PO rows found."))
        headers = set(lines[0].keys())
        missing = required - headers
        if missing:
            raise UserError(_("File is missing columns: %s") % ", ".join(sorted(missing)))
        grouped = {}
        created = 0
        updated = 0
        for mapped in lines:
            po = mapped.get("po")
            if not po:
                raise UserError(_("Each row needs a PO number."))
            try:
                ordered = float(mapped.get("ordered") or 0)
                received = float(mapped.get("received") or 0)
            except ValueError as err:
                raise UserError(_("Quantity must be numeric for PO %s.") % po) from err
            pending = float(mapped.get("pending") or (ordered - received))
            material_line = {
                "code": mapped.get("material_code") or "",
                "description": mapped.get("material_desc") or mapped.get("description") or "",
                "uom": mapped.get("uom") or "PCS",
                "ordered": ordered,
                "received": received,
                "variance": received - ordered,
                "pending": pending,
                "pipe_id": mapped.get("pipe_id") or mapped.get("pipe") or "",
                "grn_no": mapped.get("grn_no") or "",
                "spec": mapped.get("spec") or "",
                "length": mapped.get("length") or "",
                "diameter": mapped.get("diameter") or "",
                "grade": mapped.get("grade") or "",
                "batch": mapped.get("batch") or "",
                "demo_note": mapped.get("demo_note") or mapped.get("use_in_demo") or "",
            }
            invoice = mapped.get("invoice_no") or mapped.get("invoice") or ""
            if po not in grouped:
                grouped[po] = {
                    "vals": {
                        "name": po,
                        "supplier": mapped.get("supplier"),
                        "po_date": mapped.get("po_date") or mapped.get("date"),
                        "expected_delivery": mapped.get("expected_delivery") or mapped.get("eta"),
                        "plant": mapped.get("plant") or "Plant 1",
                        "status": mapped.get("status") or "Open",
                        "remarks": mapped.get("remarks") or "-",
                        "material": material_line["description"],
                        "arrival_datetime": mapped.get("arrival_datetime"),
                        "truck_no": mapped.get("truck_no") or mapped.get("truck"),
                        "transporter": mapped.get("transporter"),
                        "reference_no": mapped.get("reference_no") or mapped.get("reference") or mapped.get("grn_no") or invoice,
                        "invoice_no": invoice,
                        "notes": mapped.get("notes"),
                        "import_file": filename,
                    },
                    "materials": [],
                }
            grouped[po]["materials"].append(material_line)
        for po, pack in grouped.items():
            mats = pack["materials"]
            pack["vals"].update(
                {
                    "item_count": len(mats),
                    "ordered_qty": sum(m["ordered"] for m in mats),
                    "received_qty": sum(m["received"] for m in mats),
                    "pending_qty": sum(m["pending"] for m in mats),
                    "materials_json": json.dumps(mats),
                }
            )
            existing = self.search([("name", "=", po)], limit=1)
            if existing:
                existing.write(pack["vals"])
                updated += 1
            else:
                self.create(pack["vals"])
                created += 1
        return {
            "created": created,
            "updated": updated,
            "total": created + updated,
            "rows": self.yms_get_inbound_pos(),
            "history": self.yms_get_import_history(),
        }

    def _yms_seed_yard_rows(self, lines):
        Product = self.env["product.product"]
        Lot = self.env["stock.lot"]
        Receive = self.env["yms.stock.receive"]
        ReceiveLine = self.env["yms.stock.receive.line"]
        locs = self.env["stock.location"].yms_get_putaway_locations()
        yards = locs.get("yards") or []
        bins = locs.get("bins") or []
        plates = locs.get("plates") or []
        seeded = skipped = 0
        for mapped in lines:
            pipe_id = (mapped.get("pipe_id") or mapped.get("pipe") or "").strip()
            code = (mapped.get("material_code") or "PIPE-24").strip()
            po = (mapped.get("po") or "").strip()
            stage = (mapped.get("stage") or "located").strip().lower()
            if not pipe_id:
                continue
            product = Product.search([("default_code", "=", code)], limit=1)
            if not product:
                raise UserError(_("Material %s is not in Inventory products.") % code)
            existing = Lot.search([("name", "=", pipe_id), ("product_id", "=", product.id)], limit=1)
            if existing:
                skipped += 1
                continue
            wizard = Receive.create({"product_id": product.id, "origin": po or "DEMO"})
            ReceiveLine.create({"wizard_id": wizard.id, "pipe_id": pipe_id})
            wizard.action_receive_ui()
            lot = Lot.search([("name", "=", pipe_id), ("product_id", "=", product.id)], limit=1)
            if stage in ("qc_fail", "failed"):
                lot.yms_save_quality_check_ui(lot.id, "failed", "Demo failed visual check")
            elif stage in ("qc_hold", "hold"):
                lot.yms_save_quality_check_ui(lot.id, "hold", "Demo quality hold")
            elif stage != "pending":
                lot.yms_save_quality_check_ui(lot.id, "passed", False)
                if yards and bins and plates:
                    yard = yards[0]
                    bin_opts = [b for b in bins if b.get("yard_id") == yard["id"]] or bins
                    bin_loc = bin_opts[0]
                    plate_opts = [p for p in plates if p.get("bin_id") == bin_loc["id"]] or plates
                    plate = plate_opts[0]
                    lot.yms_putaway_ui(lot.id, yard["id"], bin_loc["id"], plate["id"])
                    if stage in ("moved", "move") and len(plate_opts) > 1:
                        other = plate_opts[1]
                        self.env["stock.picking"].yms_move_pipe_ui(
                            lot.id, yard["id"], bin_loc["id"], other["id"], 1.0, plate["id"], "Demo yard transfer"
                        )
            seeded += 1
        return {"seeded": seeded, "skipped": skipped}
